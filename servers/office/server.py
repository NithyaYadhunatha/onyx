"""Onyx office server: Excel workbooks, bank statements, SMS, and the expense sheet."""

from __future__ import annotations

import datetime as dt
import json
import os
import sys
from collections import defaultdict

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.dirname(__file__))

import openpyxl  # noqa: E402
from openpyxl.chart import BarChart, Reference  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter, range_boundaries  # noqa: E402
from mcp.server.fastmcp import FastMCP  # noqa: E402

import finance  # noqa: E402
from common import ONYX_HOME, log, reply  # noqa: E402

mcp = FastMCP(
    "onyx_office",
    log_level="WARNING",
    instructions="Read and write Excel workbooks, parse bank statements and bank SMS, and keep the user's expense sheet.",
)

DEFAULT_WORKBOOK = os.environ.get("ONYX_EXPENSES") or os.path.join(os.path.expanduser("~"), "Documents", "Onyx", "Expenses.xlsx")
PENDING = os.path.join(ONYX_HOME, "expense_pending.json")

TX_HEADERS = ["Date", "Time", "Merchant", "Category", "Type", "Amount", "Account", "Bank", "Ref", "Source", "Description", "Month", "Key", "Added"]
COL = {h: i + 1 for i, h in enumerate(TX_HEADERS)}
HEAD_FILL = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def _path(p: str | None) -> str:
    p = os.path.expandvars(os.path.expanduser(p or DEFAULT_WORKBOOK))
    return os.path.abspath(p)


def _locked_error(path: str) -> PermissionError:
    return PermissionError(f"{path} is open in another program (probably Excel). Close it and try again.")


def _cell_value(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    return v


# ------------------------------------------------------------------ generic Excel


@mcp.tool()
def excel_info(path: str) -> str:
    """Sheets in a workbook with their size and first row (usually the headers)."""
    p = _path(path)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        out.append({"sheet": ws.title, "rows": ws.max_row, "columns": ws.max_column, "first_row": [_cell_value(c) for c in first]})
    return reply({"path": p, "sheets": out})


@mcp.tool()
def excel_read(path: str, sheet: str = "", cells: str = "", max_rows: int = 200) -> str:
    """Read values (formulas give their last calculated value). `cells` like "A1:F50"; default the whole sheet."""
    p = _path(path)
    wb = openpyxl.load_workbook(p, data_only=True, read_only=not cells)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = []
    if cells:
        for row in ws[cells]:
            rows.append([_cell_value(c.value) for c in (row if isinstance(row, tuple) else (row,))])
    else:
        for row in ws.iter_rows(values_only=True):
            rows.append([_cell_value(c) for c in row])
            if len(rows) >= max_rows:
                break
    return reply({"sheet": ws.title, "rows": rows[:max_rows], "truncated": len(rows) > max_rows})


@mcp.tool()
def excel_write(path: str, rows: list[list], sheet: str = "Sheet1", start_cell: str = "A1", mode: str = "overwrite",
                header: bool = True, autofit: bool = True) -> str:
    """Write rows (a list of lists) to a workbook, creating the file and sheet if needed.
    mode: "overwrite" writes at start_cell; "append" adds below the last row. Text starting with "=" is a formula.
    header=True makes the first written row bold (overwrite mode)."""
    p = _path(path)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    if os.path.exists(p):
        wb = openpyxl.load_workbook(p)
    else:
        wb = openpyxl.Workbook()
        wb.active.title = sheet
    ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
    if mode == "append":
        r0 = ws.max_row + 1 if ws.max_row > 1 or ws.cell(1, 1).value is not None else 1
        c0 = 1
    else:
        c0, r0, _, _ = range_boundaries(start_cell + ":" + start_cell)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = ws.cell(r0 + i, c0 + j, v)
            if header and mode != "append" and i == 0:
                cell.font = HEAD_FONT
                cell.fill = HEAD_FILL
    if autofit:
        _autofit(ws)
    try:
        wb.save(p)
    except PermissionError:
        raise _locked_error(p)
    return reply({"path": p, "sheet": ws.title, "written_rows": len(rows), "from_row": r0})


def _autofit(ws, cap: int = 60):
    widths = defaultdict(int)
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths[c.column], min(cap, len(str(c.value)) + 2))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(8, w)


# ------------------------------------------------------------------ parsing


@mcp.tool()
def parse_statement(path: str, password: str = "") -> str:
    """Read transactions from a bank or UPI-app statement (PDF, CSV or XLSX: PhonePe, Google Pay,
    Paytm, HDFC, SBI, ICICI, Axis and most others). Returns them for expense_add."""
    p = _path(path)
    tx = finance.parse_statement_file(p, password or None)
    rules = finance.load_rules()
    for t in tx:
        t["category"] = finance.categorize(t, rules)
    debits = sum(t["amount"] for t in tx if t["direction"] == "debit")
    credits = sum(t["amount"] for t in tx if t["direction"] == "credit")
    note = "" if tx else "No transactions recognised. Check with excel_read (CSV/XLSX), or the PDF may be scanned or protected."
    return reply({"file": p, "count": len(tx), "total_debit": round(debits, 2), "total_credit": round(credits, 2), "transactions": tx, "note": note})


@mcp.tool()
def parse_sms(messages: list[dict]) -> str:
    """Turn bank SMS into transactions. Each message: {address (sender), body, date_ms}.
    Skips OTPs, offers, payment requests and failed payments."""
    tx = []
    rules = finance.load_rules()
    for m in messages:
        t = finance.parse_sms(str(m.get("address") or m.get("from") or ""), str(m.get("body") or ""), m.get("date_ms"))
        if t:
            t["category"] = finance.categorize(t, rules)
            tx.append(t)
    return reply({"messages": len(messages), "count": len(tx), "transactions": tx})


# ------------------------------------------------------------------ the expense workbook


def _open_expenses(p: str):
    os.makedirs(os.path.dirname(p), exist_ok=True)
    if os.path.exists(p):
        wb = openpyxl.load_workbook(p)
    else:
        wb = openpyxl.Workbook()
        wb.active.title = "Transactions"
    if "Transactions" not in wb.sheetnames:
        wb.create_sheet("Transactions", 0)
    ws = wb["Transactions"]
    if ws.cell(1, 1).value != "Date":
        for i, h in enumerate(TX_HEADERS, 1):
            c = ws.cell(1, i, h)
            c.font, c.fill = HEAD_FONT, HEAD_FILL
        ws.freeze_panes = "A2"
        widths = {"Date": 12, "Time": 8, "Merchant": 26, "Category": 18, "Type": 8, "Amount": 12, "Account": 9, "Bank": 12,
                  "Ref": 16, "Source": 18, "Description": 60, "Month": 9, "Key": 6, "Added": 18}
        for h, w in widths.items():
            ws.column_dimensions[get_column_letter(COL[h])].width = w
        ws.column_dimensions[get_column_letter(COL["Key"])].hidden = True
    return wb, ws


def _rows(ws) -> list[dict]:
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or r[0] is None:
            continue
        d = {h: (r[i] if i < len(r) else None) for i, h in enumerate(TX_HEADERS)}
        if isinstance(d["Date"], (dt.date, dt.datetime)):
            d["Date"] = d["Date"].strftime("%Y-%m-%d")
        out.append(d)
    return out


def _month_totals(rows: list[dict]):
    spent = defaultdict(lambda: defaultdict(float))
    earned = defaultdict(float)
    for r in rows:
        month = str(r["Month"] or str(r["Date"])[:7])
        amt = float(r["Amount"] or 0)
        if r["Type"] == "Debit":
            spent[month][r["Category"] or "Other"] += amt
        else:
            earned[month] += amt
    return spent, earned


def _rebuild_summary(wb, rows: list[dict]):
    """One row per month, one column per category, as live SUMIFS formulas so
    edits to Transactions (a fixed category, a deleted row) flow through."""
    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    ws = wb.create_sheet("Summary", 1)
    spent, earned = _month_totals(rows)
    months = sorted(set(spent) | set(earned))
    cats = sorted({c for m in spent.values() for c in m}, key=lambda c: -sum(spent[m].get(c, 0) for m in spent))
    head = ["Month", "Total spent", "Money in", *cats]
    for i, h in enumerate(head, 1):
        c = ws.cell(1, i, h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    T = "Transactions"
    amt = f"{T}!${get_column_letter(COL['Amount'])}:${get_column_letter(COL['Amount'])}"
    mon = f"{T}!${get_column_letter(COL['Month'])}:${get_column_letter(COL['Month'])}"
    typ = f"{T}!${get_column_letter(COL['Type'])}:${get_column_letter(COL['Type'])}"
    cat = f"{T}!${get_column_letter(COL['Category'])}:${get_column_letter(COL['Category'])}"
    for r, m in enumerate(months, 2):
        ws.cell(r, 1, m)
        ws.cell(r, 2, f'=SUMIFS({amt},{mon},$A{r},{typ},"Debit")')
        ws.cell(r, 3, f'=SUMIFS({amt},{mon},$A{r},{typ},"Credit")')
        for j, c in enumerate(cats, 4):
            ws.cell(r, j, f'=SUMIFS({amt},{mon},$A{r},{typ},"Debit",{cat},{get_column_letter(j)}$1)')
        for j in range(2, len(head) + 1):
            ws.cell(r, j).number_format = "₹#,##0"
    ws.column_dimensions["A"].width = 10
    for j in range(2, len(head) + 1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, len(head[j - 1]) + 3)
    ws.freeze_panes = "B2"
    if months:
        chart = BarChart()
        chart.title = "Spending by month"
        chart.y_axis.title = "₹"
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(months) + 1), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(months) + 1))
        chart.height, chart.width = 8, 18
        ws.add_chart(chart, f"A{len(months) + 4}")


def _load_pending() -> list[dict]:
    try:
        with open(PENDING, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return []


def _save_pending(tx: list[dict]):
    os.makedirs(os.path.dirname(PENDING), exist_ok=True)
    with open(PENDING, "w", encoding="utf-8") as f:
        json.dump(tx, f, ensure_ascii=False)


@mcp.tool()
def expense_add(transactions: list[dict], workbook: str = "") -> str:
    """Add transactions (from parse_statement or parse_sms) to the expense workbook. Skips duplicates
    already in it, categorises, and rebuilds the monthly Summary sheet and chart.
    If the workbook is open in Excel, they are queued and added next time."""
    p = _path(workbook)
    tx = _load_pending() + list(transactions)
    try:
        wb, ws = _open_expenses(p)
    except PermissionError:
        _save_pending(tx)
        raise _locked_error(p)
    rows = _rows(ws)
    keys = {r["Key"] for r in rows if r["Key"]}
    rules = finance.load_rules()
    added, added_spent, dupes = 0, 0.0, 0
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    for t in tx:
        try:
            amount = round(float(t.get("amount") or 0), 2)
        except (TypeError, ValueError):
            continue
        date = finance.parse_date(t.get("date")) or dt.date.today().isoformat()
        if not amount:
            continue
        t = {**t, "date": date, "amount": amount}
        key = finance.tx_key(t)
        if key in keys:
            dupes += 1
            continue
        keys.add(key)
        category = t.get("category") or finance.categorize(t, rules)
        kind = "Debit" if str(t.get("direction", "debit")).lower().startswith("d") else "Credit"
        row = {
            "Date": date, "Time": t.get("time", ""), "Merchant": t.get("merchant", ""), "Category": category, "Type": kind,
            "Amount": amount, "Account": t.get("account", ""), "Bank": t.get("bank", ""), "Ref": t.get("ref", ""),
            "Source": t.get("source", ""), "Description": t.get("description", ""), "Month": date[:7], "Key": key, "Added": now,
        }
        ws.append([row[h] for h in TX_HEADERS])
        ws.cell(ws.max_row, COL["Amount"]).number_format = "₹#,##0.00"
        rows.append(row)
        added += 1
        if kind == "Debit":
            added_spent += amount
    # Keep the sheet in date order so it reads like a statement.
    if added:
        data = sorted(_rows(ws), key=lambda r: (str(r["Date"]), str(r["Time"] or "")))
        ws.delete_rows(2, ws.max_row)
        for r in data:
            ws.append([r[h] for h in TX_HEADERS])
            ws.cell(ws.max_row, COL["Amount"]).number_format = "₹#,##0.00"
        rows = data
    _rebuild_summary(wb, rows)
    try:
        wb.save(p)
    except PermissionError:
        _save_pending(tx)
        raise _locked_error(p)
    if os.path.exists(PENDING):
        os.remove(PENDING)
    spent, earned = _month_totals(rows)
    this_month = dt.date.today().strftime("%Y-%m")
    return reply({
        "workbook": p, "added": added, "duplicates_skipped": dupes, "added_spent": round(added_spent, 2),
        "month": this_month, "month_spent": round(sum(spent.get(this_month, {}).values()), 2),
        "month_top": sorted(((c, round(v, 2)) for c, v in spent.get(this_month, {}).items()), key=lambda x: -x[1])[:5],
    })


@mcp.tool()
def expense_summary(month: str = "", workbook: str = "") -> str:
    """Spending for a month ("2026-10"; default this month): total, by category, top merchants, money in.
    month="all" gives every month's total."""
    p = _path(workbook)
    if not os.path.exists(p):
        return reply({"workbook": p, "note": "No expense workbook yet. Add a statement or connect the phone to start."})
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    rows = _rows(wb["Transactions"])
    spent, earned = _month_totals(rows)
    if month == "all":
        return reply({"workbook": p, "months": {m: round(sum(c.values()), 2) for m, c in sorted(spent.items())}})
    m = month or dt.date.today().strftime("%Y-%m")
    merchants = defaultdict(float)
    for r in rows:
        if str(r["Month"]) == m and r["Type"] == "Debit":
            merchants[r["Merchant"] or "Unknown"] += float(r["Amount"] or 0)
    return reply({
        "workbook": p, "month": m, "total_spent": round(sum(spent.get(m, {}).values()), 2),
        "money_in": round(earned.get(m, 0), 2),
        "by_category": dict(sorted(((c, round(v, 2)) for c, v in spent.get(m, {}).items()), key=lambda x: -x[1])),
        "top_merchants": sorted(((k, round(v, 2)) for k, v in merchants.items()), key=lambda x: -x[1])[:8],
        "transactions": sum(1 for r in rows if str(r["Month"]) == m),
    })


@mcp.tool()
def set_category_rule(keyword: str, category: str, recategorize: bool = True, workbook: str = "") -> str:
    """Always put transactions whose merchant or description contains `keyword` into `category`
    (e.g. "chai point" -> "Food & Dining"). recategorize=True also fixes existing rows."""
    rules = finance.save_rule(keyword, category)
    changed = 0
    p = _path(workbook)
    if recategorize and os.path.exists(p):
        wb, ws = _open_expenses(p)
        kw = keyword.lower()
        for r in range(2, ws.max_row + 1):
            hay = f"{ws.cell(r, COL['Merchant']).value or ''} {ws.cell(r, COL['Description']).value or ''}".lower()
            if kw in hay and ws.cell(r, COL["Type"]).value == "Debit":
                ws.cell(r, COL["Category"], category)
                changed += 1
        _rebuild_summary(wb, _rows(ws))
        try:
            wb.save(p)
        except PermissionError:
            raise _locked_error(p)
    return reply({"rules": len(rules), "recategorized": changed})


if __name__ == "__main__":
    log("[office] Onyx office server ready")
    mcp.run()
